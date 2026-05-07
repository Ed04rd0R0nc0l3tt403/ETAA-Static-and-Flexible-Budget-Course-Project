import pandas as pd
from src.costs_bdg import compute_production_cost_bdg

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def compute_income_bdg(
    sales_filepath:      str = "data/bic_sales_data.csv",
    production_filepath: str = "data/bic_production_data.csv",
    cost_filepath:       str = "data/bic_cost_drivers.csv",
    opex_filepath:       str = "data/operating_expenses.csv"
):
    opex_df = pd.read_csv(opex_filepath).set_index("segment")

    raw_sales = pd.read_csv(sales_filepath)
    raw_sales["revenue"] = raw_sales["units_millions"] * raw_sales["unit_price_eur"]
    sales_q = (
        raw_sales
        .groupby(["year", "quarter", "segment"])["revenue"]
        .sum()
        .reset_index()
    )

    cost_q, _ = compute_production_cost_bdg(production_filepath, sales_filepath, cost_filepath)

    df = sales_q.merge(
        cost_q[["year", "quarter", "segment", "total_cost"]],
        on=["year", "quarter", "segment"],
        how="left"
    )

    df["gross_profit"]        = df["revenue"] - df["total_cost"]
    df["gross_profit_margin"] = df["gross_profit"] / df["revenue"]

    df["marketing"]  = df["revenue"] * df.apply(lambda r: opex_df.loc[r["segment"], "marketing_rate_revenue"], axis=1)
    df["rd"]         = df["revenue"] * df.apply(lambda r: opex_df.loc[r["segment"], "rd_rate_revenue"], axis=1)
    df["admin"]      = df.apply(lambda r: opex_df.loc[r["segment"], "admin_fixed_monthly"] * 3 / 1_000_000, axis=1)
    df["total_opex"] = df["marketing"] + df["rd"] + df["admin"]

    df["ebit"]        = df["gross_profit"] - df["total_opex"]
    df["ebit_margin"] = df["ebit"] / df["revenue"]

    quarterly_df = df.sort_values(["year", "quarter", "segment"]).reset_index(drop=True)

    yearly_df = (
        quarterly_df
        .groupby(["year", "segment"])
        .agg(
            revenue      = ("revenue",      "sum"),
            total_cost   = ("total_cost",   "sum"),
            gross_profit = ("gross_profit", "sum"),
            marketing    = ("marketing",    "sum"),
            rd           = ("rd",           "sum"),
            admin        = ("admin",        "sum"),
            total_opex   = ("total_opex",   "sum"),
            ebit         = ("ebit",         "sum"),
        )
        .assign(
            gross_profit_margin = lambda d: d["gross_profit"] / d["revenue"],
            ebit_margin         = lambda d: d["ebit"]         / d["revenue"],
        )
        .reset_index()
    )

    return quarterly_df, yearly_df


def export_income_bdg_excel(
    sales_filepath:      str = "data/bic_sales_data.csv",
    production_filepath: str = "data/bic_production_data.csv",
    cost_filepath:       str = "data/bic_cost_drivers.csv",
    opex_filepath:       str = "data/operating_expenses.csv",
    output_path:         str = "output/sales_budget.xlsx"
):
    quarterly_df, yearly_df = compute_income_bdg(
        sales_filepath, production_filepath, cost_filepath, opex_filepath
    )

    wb = load_workbook(output_path)

    if "Income Budget" in wb.sheetnames:
        del wb["Income Budget"]
    ws = wb.create_sheet("Income Budget")

    header_font  = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", start_color="1F4E79")
    section_fill = PatternFill("solid", start_color="2E75B6")
    total_fill   = PatternFill("solid", start_color="D6E4F0")
    total_font   = Font(name="Arial", bold=True, color="000000")
    center       = Alignment(horizontal="center")
    thin         = Side(style="thin", color="CCCCCC")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_cell(cell, font=None, fill=None, alignment=None, number_format=None):
        if font:          cell.font          = font
        if fill:          cell.fill          = fill
        if alignment:     cell.alignment     = alignment
        if number_format: cell.number_format = number_format
        cell.border = border

    ws.merge_cells("A1:M1")
    ws["A1"] = "BIC Corporation — Budgeted Income Statement"
    ws["A1"].font      = Font(name="Arial", bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    # --- Quarterly Detail ---
    ws.append([])
    ws.append(["QUARTERLY DETAIL"] + [""] * 12)
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    q_headers = ["Year", "Quarter", "Segment", "Revenue (€M)", "Total Cost (€M)",
                 "Gross Profit (€M)", "GP Margin %", "Marketing (€M)", "R&D (€M)",
                 "Admin (€M)", "Total OpEx (€M)", "EBIT (€M)", "EBIT Margin %"]
    ws.append(q_headers)
    for col in range(1, len(q_headers) + 1):
        style_cell(ws.cell(ws.max_row, col), font=header_font, fill=header_fill, alignment=center)

    for _, row in quarterly_df.sort_values(["year", "quarter", "segment"]).iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value  = int(row["year"])
        ws.cell(r, 2).value  = row["quarter"]
        ws.cell(r, 3).value  = row["segment"]
        ws.cell(r, 4).value  = round(row["revenue"],             4)
        ws.cell(r, 5).value  = round(row["total_cost"],          4)
        ws.cell(r, 6).value  = round(row["gross_profit"],        4)
        ws.cell(r, 7).value  = round(row["gross_profit_margin"], 4)
        ws.cell(r, 8).value  = round(row["marketing"],           4)
        ws.cell(r, 9).value  = round(row["rd"],                  4)
        ws.cell(r, 10).value = round(row["admin"],               4)
        ws.cell(r, 11).value = round(row["total_opex"],          4)
        ws.cell(r, 12).value = round(row["ebit"],                4)
        ws.cell(r, 13).value = round(row["ebit_margin"],         4)
        for col in range(1, 14):
            fmt = '0.00%' if col in [7, 13] else '#,##0.0000;(#,##0.0000);-'
            style_cell(ws.cell(r, col), number_format=fmt)

    # --- Yearly Summary ---
    ws.append([])
    ws.append(["YEARLY SUMMARY"] + [""] * 11)
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    y_headers = ["Year", "Segment", "Revenue (€M)", "Total Cost (€M)",
                 "Gross Profit (€M)", "GP Margin %", "Marketing (€M)", "R&D (€M)",
                 "Admin (€M)", "Total OpEx (€M)", "EBIT (€M)", "EBIT Margin %"]
    ws.append(y_headers)
    for col in range(1, len(y_headers) + 1):
        style_cell(ws.cell(ws.max_row, col), font=header_font, fill=header_fill, alignment=center)

    for _, row in yearly_df.sort_values(["year", "segment"]).iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value  = int(row["year"])
        ws.cell(r, 2).value  = row["segment"]
        ws.cell(r, 3).value  = round(row["revenue"],             4)
        ws.cell(r, 4).value  = round(row["total_cost"],          4)
        ws.cell(r, 5).value  = round(row["gross_profit"],        4)
        ws.cell(r, 6).value  = round(row["gross_profit_margin"], 4)
        ws.cell(r, 7).value  = round(row["marketing"],           4)
        ws.cell(r, 8).value  = round(row["rd"],                  4)
        ws.cell(r, 9).value  = round(row["admin"],               4)
        ws.cell(r, 10).value = round(row["total_opex"],          4)
        ws.cell(r, 11).value = round(row["ebit"],                4)
        ws.cell(r, 12).value = round(row["ebit_margin"],         4)
        for col in range(1, 13):
            fmt = '0.00%' if col in [6, 12] else '#,##0.0000;(#,##0.0000);-'
            style_cell(ws.cell(r, col), number_format=fmt)

    # --- Yearly Totals ---
    ws.append([])
    ws.append(["YEARLY TOTALS"] + [""] * 5)
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    t_headers = ["Year", "Total Revenue (€M)", "Total Cost (€M)",
                 "Total Gross Profit (€M)", "Total OpEx (€M)", "Total EBIT (€M)"]
    ws.append(t_headers)
    for col in range(1, len(t_headers) + 1):
        style_cell(ws.cell(ws.max_row, col), font=header_font, fill=header_fill, alignment=center)

    for year in sorted(yearly_df["year"].unique()):
        sub = yearly_df[yearly_df["year"] == year]
        r   = ws.max_row + 1
        ws.cell(r, 1).value = int(year)
        ws.cell(r, 2).value = round(sub["revenue"].sum(),      4)
        ws.cell(r, 3).value = round(sub["total_cost"].sum(),   4)
        ws.cell(r, 4).value = round(sub["gross_profit"].sum(), 4)
        ws.cell(r, 5).value = round(sub["total_opex"].sum(),   4)
        ws.cell(r, 6).value = round(sub["ebit"].sum(),         4)
        for col in range(1, 7):
            style_cell(ws.cell(r, col), font=total_font, fill=total_fill,
                       number_format='#,##0.0000;(#,##0.0000);-')

    col_widths = [8, 10, 22, 18, 20, 20, 12, 18, 12, 12, 18, 14, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(output_path)
    print(f"Income budget exported to '{output_path}' — sheet: Income Budget")