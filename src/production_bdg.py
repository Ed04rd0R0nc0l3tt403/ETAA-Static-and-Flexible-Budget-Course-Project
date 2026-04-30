import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def compute_production_bdg(
    sales_filepath: str = "data/bic_sales_data.csv",
    production_filepath: str = "data/bic_production_data.csv"
):
    sales_df = pd.read_csv(sales_filepath).sort_values(["segment", "year", "month_num"]).reset_index(drop=True)
    rates_df = pd.read_csv(production_filepath).set_index("segment")

    # Last month of each quarter → last month of next quarter
    quarter_end_to_next = {3: 6, 6: 9, 9: 12, 12: 3}
    quarter_end_months = set(quarter_end_to_next.keys())

    results = []

    for segment, group in sales_df.groupby("segment"):
        group = group.sort_values(["year", "month_num"]).reset_index(drop=True)
        end_inv_rate = rates_df.loc[segment, "end_inv_rate"]
        beg_inv_rate = rates_df.loc[segment, "beg_inv_rate"]

        def get_qty(year, month):
            """Lookup quantity for a specific year/month within this segment."""
            match = group[(group["year"] == year) & (group["month_num"] == month)]
            return float(match["units_millions"].values[0]) if not match.empty else None

        # --- Ending Inventory ---
        end_inv_list = []
        for _, row in group.iterrows():
            month = int(row["month_num"])
            year  = int(row["year"])

            if month in quarter_end_months:
                # Last month of quarter → use last month of next quarter
                next_end_month = quarter_end_to_next[month]
                next_year = year + 1 if month == 12 else year
                qty = get_qty(next_year, next_end_month)
            else:
                # Regular month → next month
                qty = get_qty(year, month + 1)

            # Fallback: if next month has no data for this segment, use current qty
            end_inv_list.append((qty if qty is not None else row["units_millions"]) * end_inv_rate)

        group["end_inv"] = end_inv_list

        # --- Beginning Inventory ---
        beg_inv_list = []
        for i in range(len(group)):
            if i == 0:
                beg_inv_list.append(group.iloc[0]["units_millions"] * beg_inv_rate)
            else:
                beg_inv_list.append(group.iloc[i - 1]["end_inv"])   # prev month end_inv

        group["beg_inv"] = beg_inv_list

        # --- Units to Produce ---
        group["units_to_produce"] = group["units_millions"] + group["end_inv"] - group["beg_inv"]

        results.append(group)

    monthly_df = (
        pd.concat(results)
        .sort_values(["year", "month_num", "segment"])
        .reset_index(drop=True)
    )

    # --- Quarterly Aggregation ---
    quarterly_rows = []
    for (year, quarter, segment), g in monthly_df.groupby(["year", "quarter", "segment"]):
        g = g.sort_values("month_num")
        quarterly_rows.append({
            "year":             year,
            "quarter":          quarter,
            "segment":          segment,
            "units_sold":       g["units_millions"].sum(),
            "beg_inv":          g["beg_inv"].iloc[0],    # first month of quarter
            "end_inv":          g["end_inv"].iloc[-1],   # last month of quarter
            "units_to_produce": g["units_millions"].sum() + g["end_inv"].iloc[-1] - g["beg_inv"].iloc[0],
        })

    quarterly_df = (
        pd.DataFrame(quarterly_rows)
        .sort_values(["year", "quarter", "segment"])
        .reset_index(drop=True)
    )

    return monthly_df, quarterly_df

def export_production_bdg_excel(
    sales_filepath:      str = "data/bic_sales_data.csv",
    production_filepath: str = "data/bic_production_data.csv",
    output_path:         str = "output/sales_budget.xlsx"
):
    _, quarterly = compute_production_bdg(sales_filepath, production_filepath)

    # yearly aggregation
    yearly = (
        quarterly
        .groupby(["year", "segment"])
        .agg(
            units_sold       = ("units_sold",       "sum"),
            beg_inv          = ("beg_inv",           "first"),
            end_inv          = ("end_inv",           "last"),
            units_to_produce = ("units_to_produce",  "sum")
        )
        .reset_index()
    )

    # --- load existing workbook ---
    wb = load_workbook(output_path)

    # --- remove sheet if already exists (avoids duplicates on re-run) ---
    if "Production Budget" in wb.sheetnames:
        del wb["Production Budget"]

    ws = wb.create_sheet("Production Budget")

    # --- Styles ---
    header_font  = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", start_color="1F4E79")
    section_fill = PatternFill("solid", start_color="2E75B6")
    total_fill   = PatternFill("solid", start_color="D6E4F0")
    total_font   = Font(name="Arial", bold=True, color="000000")
    center       = Alignment(horizontal="center")
    thin         = Side(style="thin", color="CCCCCC")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_cell(cell, font=None, fill=None, alignment=None, number_format=None):
        if font:          cell.font          = font
        if fill:          cell.fill          = fill
        if alignment:     cell.alignment     = alignment
        if number_format: cell.number_format = number_format
        cell.border = border

    # --- Title ---
    ws.merge_cells("A1:F1")
    ws["A1"] = "BIC Corporation — Production Budget"
    ws["A1"].font      = Font(name="Arial", bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    # ------------------------------------------------------------------ #
    #  SECTION 1 — QUARTERLY DETAIL                                        #
    # ------------------------------------------------------------------ #
    ws.append([])
    ws.append(["QUARTERLY DETAIL", "", "", "", "", ""])
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    q_headers = ["Year", "Quarter", "Segment",
                 "Units Sold (M)", "Beg. Inventory (M)",
                 "End. Inventory (M)", "Units to Produce (M)"]
    ws.append(q_headers)
    for col in range(1, len(q_headers) + 1):
        style_cell(ws.cell(ws.max_row, col),
                   font=header_font, fill=header_fill, alignment=center)

    q_sorted = quarterly.sort_values(["year", "quarter", "segment"])
    for _, row in q_sorted.iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = row["quarter"]
        ws.cell(r, 3).value = row["segment"]
        ws.cell(r, 4).value = round(row["units_sold"],       4)
        ws.cell(r, 5).value = round(row["beg_inv"],          4)
        ws.cell(r, 6).value = round(row["end_inv"],          4)
        ws.cell(r, 7).value = round(row["units_to_produce"], 4)
        for col in range(1, 8):
            style_cell(ws.cell(r, col),
                       number_format='#,##0.0000;(#,##0.0000);-')

    # ------------------------------------------------------------------ #
    #  SECTION 2 — YEARLY SUMMARY                                          #
    # ------------------------------------------------------------------ #
    ws.append([])
    ws.append(["YEARLY SUMMARY", "", "", "", "", ""])
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    y_headers = ["Year", "Segment", "Units Sold (M)",
                 "Beg. Inventory (M)", "End. Inventory (M)",
                 "Units to Produce (M)"]
    ws.append(y_headers)
    for col in range(1, len(y_headers) + 1):
        style_cell(ws.cell(ws.max_row, col),
                   font=header_font, fill=header_fill, alignment=center)

    y_sorted = yearly.sort_values(["year", "segment"])
    for _, row in y_sorted.iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = row["segment"]
        ws.cell(r, 3).value = round(row["units_sold"],       4)
        ws.cell(r, 4).value = round(row["beg_inv"],          4)
        ws.cell(r, 5).value = round(row["end_inv"],          4)
        ws.cell(r, 6).value = round(row["units_to_produce"], 4)
        for col in range(1, 7):
            style_cell(ws.cell(r, col),
                       number_format='#,##0.0000;(#,##0.0000);-')

    # ------------------------------------------------------------------ #
    #  SECTION 3 — YEARLY TOTALS ACROSS SEGMENTS                           #
    # ------------------------------------------------------------------ #
    ws.append([])
    ws.append(["YEARLY TOTALS", "", "", "", "", ""])
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    totals_headers = ["Year", "Total Units Sold (M)", "Total Units to Produce (M)"]
    ws.append(totals_headers)
    for col in range(1, len(totals_headers) + 1):
        style_cell(ws.cell(ws.max_row, col),
                   font=header_font, fill=header_fill, alignment=center)

    for year in sorted(yearly["year"].unique()):
        sub = yearly[yearly["year"] == year]
        r   = ws.max_row + 1
        ws.cell(r, 1).value = int(year)
        ws.cell(r, 2).value = round(sub["units_sold"].sum(),       4)
        ws.cell(r, 3).value = round(sub["units_to_produce"].sum(), 4)
        for col in range(1, 4):
            style_cell(ws.cell(r, col),
                       font=total_font, fill=total_fill,
                       number_format='#,##0.0000;(#,##0.0000);-')

    # --- Column widths ---
    col_widths = [8, 10, 22, 18, 20, 22, 22]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(output_path)
    print(f"Production budget exported to '{output_path}' — sheet: Production Budget")