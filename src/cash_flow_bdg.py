import pandas as pd
from src.costs_bdg import compute_production_cost_bdg

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def compute_cash_flow_bdg(
    sales_filepath:      str = "data/bic_sales_data.csv",
    production_filepath: str = "data/bic_production_data.csv",
    cost_filepath:       str = "data/bic_cost_drivers.csv",
    policies_filepath:   str = "data/bic_cash_policies.csv"
):
    # --- Load data ---
    policies_df = pd.read_csv(policies_filepath)

    raw_sales = pd.read_csv(sales_filepath)
    raw_sales["revenue_eur_millions"] = raw_sales["units_millions"] * raw_sales["unit_price_eur"]
    sales_df = (
        raw_sales
        .groupby(["year", "quarter", "segment"])["revenue_eur_millions"]
        .sum()
        .reset_index()
    )

    cost_q, _ = compute_production_cost_bdg(production_filepath, sales_filepath, cost_filepath)

    # Split policies by type
    collection = policies_df[policies_df["policy_type"] == "Collection"].set_index("category")
    payment    = policies_df[policies_df["policy_type"] == "Payment"].set_index("category")

    def offset_quarter(q, y, offset):
        """Return (quarter_label, year) shifted by offset quarters."""
        new_q = q + offset
        extra_years, new_q = divmod(new_q - 1, 4)
        return f"Q{new_q + 1}", y + extra_years

    # ------------------------------------------------------------------ #
    #  CASH COLLECTIONS                                                    #
    # ------------------------------------------------------------------ #
    collection_rows = []

    for _, row in sales_df.iterrows():
        seg     = row["segment"]
        year    = int(row["year"])
        q_num   = int(row["quarter"][1])
        revenue = row["revenue_eur_millions"]
        p       = collection.loc[seg]

        q0, y0 = offset_quarter(q_num, year, 0)
        q1, y1 = offset_quarter(q_num, year, 1)
        q2, y2 = offset_quarter(q_num, year, 2)

        collection_rows += [
            {"year": y0, "quarter": q0, "segment": seg, "cash_inflow": revenue * p["current_month_pct"]},
            {"year": y1, "quarter": q1, "segment": seg, "cash_inflow": revenue * p["following_month_pct"]},
            {"year": y2, "quarter": q2, "segment": seg, "cash_inflow": revenue * p["two_months_out_pct"]},
        ]

    collections_df = (
        pd.DataFrame(collection_rows)
        .groupby(["year", "quarter", "segment"])["cash_inflow"]
        .sum()
        .reset_index()
    )

    # ------------------------------------------------------------------ #
    #  CASH PAYMENTS                                                       #
    # ------------------------------------------------------------------ #
    payment_rows = []

    for _, row in cost_q.iterrows():
        year  = int(row["year"])
        q_num = int(row["quarter"][1])
        seg   = row["segment"]

        mat_p = payment.loc["Direct Materials"]
        loh_p = payment.loc["Labor & Overhead"]

        q0, y0 = offset_quarter(q_num, year, 0)
        q1, y1 = offset_quarter(q_num, year, 1)

        mat = row["material_cost"]
        loh = row["labour_cost"] + row["total_overhead"]

        payment_rows += [
            {"year": y0, "quarter": q0, "segment": seg, "category": "Direct Materials",
             "cash_outflow": mat * mat_p["current_month_pct"]},
            {"year": y1, "quarter": q1, "segment": seg, "category": "Direct Materials",
             "cash_outflow": mat * mat_p["following_month_pct"]},
            {"year": y0, "quarter": q0, "segment": seg, "category": "Labor & Overhead",
             "cash_outflow": loh * loh_p["current_month_pct"]},
        ]

    payments_df = (
        pd.DataFrame(payment_rows)
        .groupby(["year", "quarter", "segment", "category"])["cash_outflow"]
        .sum()
        .reset_index()
    )

    # ------------------------------------------------------------------ #
    #  NET CASH POSITION                                                   #
    # ------------------------------------------------------------------ #
    total_inflows  = collections_df.groupby(["year", "quarter"])["cash_inflow"].sum().reset_index()
    total_outflows = payments_df.groupby(["year", "quarter"])["cash_outflow"].sum().reset_index()

    net_cash = (
        total_inflows
        .merge(total_outflows, on=["year", "quarter"], how="outer")
        .fillna(0)
        .assign(net_cash=lambda df: df["cash_inflow"] - df["cash_outflow"])
        .sort_values(["year", "quarter"])
        .reset_index(drop=True)
    )

    return collections_df, payments_df, net_cash

def export_cash_flow_bdg_excel(
    sales_filepath:      str = "data/bic_sales_data.csv",
    production_filepath: str = "data/bic_production_data.csv",
    cost_filepath:       str = "data/bic_cost_drivers.csv",
    policies_filepath:   str = "data/bic_cash_policies.csv",
    output_path:         str = "output/sales_budget.xlsx"
):
    collections_df, payments_df, net_cash = compute_cash_flow_bdg(
        sales_filepath, production_filepath, cost_filepath, policies_filepath
    )

    wb = load_workbook(output_path)

    if "Cash Flow Budget" in wb.sheetnames:
        del wb["Cash Flow Budget"]
    ws = wb.create_sheet("Cash Flow Budget")

    # --- Styles ---
    header_font  = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", start_color="1F4E79")
    section_fill = PatternFill("solid", start_color="2E75B6")
    total_fill   = PatternFill("solid", start_color="D6E4F0")
    total_font   = Font(name="Arial", bold=True, color="000000")
    center       = Alignment(horizontal="center")
    thin         = Side(style="thin", color="CCCCCC")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_cell(cell, font=None, fill=None, alignment=None, number_format=None):
        if font:          cell.font          = font
        if fill:          cell.fill          = fill
        if alignment:     cell.alignment     = alignment
        if number_format: cell.number_format = number_format
        cell.border = border

    # --- Title ---
    ws.merge_cells("A1:F1")
    ws["A1"] = "BIC Corporation — Cash Flow Budget"
    ws["A1"].font      = Font(name="Arial", bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    # ------------------------------------------------------------------ #
    #  SECTION 1 — CASH COLLECTIONS BY SEGMENT                            #
    # ------------------------------------------------------------------ #
    ws.append([])
    ws.append(["CASH COLLECTIONS BY SEGMENT"] + [""] * 3)
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    c_headers = ["Year", "Quarter", "Segment", "Cash Inflow (€M)"]
    ws.append(c_headers)
    for col in range(1, len(c_headers) + 1):
        style_cell(ws.cell(ws.max_row, col),
                   font=header_font, fill=header_fill, alignment=center)

    c_sorted = collections_df.sort_values(["year", "quarter", "segment"])
    for _, row in c_sorted.iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = row["quarter"]
        ws.cell(r, 3).value = row["segment"]
        ws.cell(r, 4).value = round(row["cash_inflow"], 4)
        for col in range(1, 5):
            style_cell(ws.cell(r, col),
                       number_format='#,##0.0000;(#,##0.0000);-')

    # quarterly totals
    ws.append([])
    ws.append(["COLLECTIONS QUARTERLY TOTAL"] + [""] * 2)
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    ct_headers = ["Year", "Quarter", "Total Cash Inflow (€M)"]
    ws.append(ct_headers)
    for col in range(1, len(ct_headers) + 1):
        style_cell(ws.cell(ws.max_row, col),
                   font=header_font, fill=header_fill, alignment=center)

    coll_totals = (
        collections_df.groupby(["year", "quarter"])["cash_inflow"]
        .sum().reset_index().sort_values(["year", "quarter"])
    )
    for _, row in coll_totals.iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = row["quarter"]
        ws.cell(r, 3).value = round(row["cash_inflow"], 4)
        for col in range(1, 4):
            style_cell(ws.cell(r, col),
                       number_format='#,##0.0000;(#,##0.0000);-')

    # ------------------------------------------------------------------ #
    #  SECTION 2 — CASH PAYMENTS BY SEGMENT AND CATEGORY                  #
    # ------------------------------------------------------------------ #
    ws.append([])
    ws.append(["CASH PAYMENTS BY SEGMENT & CATEGORY"] + [""] * 4)
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    p_headers = ["Year", "Quarter", "Segment", "Category", "Cash Outflow (€M)"]
    ws.append(p_headers)
    for col in range(1, len(p_headers) + 1):
        style_cell(ws.cell(ws.max_row, col),
                   font=header_font, fill=header_fill, alignment=center)

    p_sorted = payments_df.sort_values(["year", "quarter", "segment", "category"])
    for _, row in p_sorted.iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = row["quarter"]
        ws.cell(r, 3).value = row["segment"]
        ws.cell(r, 4).value = row["category"]
        ws.cell(r, 5).value = round(row["cash_outflow"], 4)
        for col in range(1, 6):
            style_cell(ws.cell(r, col),
                       number_format='#,##0.0000;(#,##0.0000);-')

    # ------------------------------------------------------------------ #
    #  SECTION 3 — NET CASH POSITION                                       #
    # ------------------------------------------------------------------ #
    ws.append([])
    ws.append(["NET CASH POSITION"] + [""] * 3)
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    n_headers = ["Year", "Quarter", "Cash Inflow (€M)", "Cash Outflow (€M)", "Net Cash (€M)"]
    ws.append(n_headers)
    for col in range(1, len(n_headers) + 1):
        style_cell(ws.cell(ws.max_row, col),
                   font=header_font, fill=header_fill, alignment=center)

    for _, row in net_cash.iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = row["quarter"]
        ws.cell(r, 3).value = round(row["cash_inflow"],  4)
        ws.cell(r, 4).value = round(row["cash_outflow"], 4)
        ws.cell(r, 5).value = round(row["net_cash"],     4)
        for col in range(1, 6):
            style_cell(ws.cell(r, col),
                       number_format='#,##0.0000;(#,##0.0000);-')

    # --- Grand totals ---
    ws.append([])
    r = ws.max_row + 1
    ws.cell(r, 1).value = "TOTAL (ALL YEARS)"
    ws.cell(r, 3).value = round(net_cash["cash_inflow"].sum(),  4)
    ws.cell(r, 4).value = round(net_cash["cash_outflow"].sum(), 4)
    ws.cell(r, 5).value = round(net_cash["net_cash"].sum(),     4)
    for col in [1, 3, 4, 5]:
        style_cell(ws.cell(r, col),
                   font=total_font, fill=total_fill,
                   number_format='#,##0.0000;(#,##0.0000);-')

    # --- Column widths ---
    col_widths = [8, 10, 22, 22, 20]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(output_path)
    print(f"Cash flow budget exported to '{output_path}' — sheet: Cash Flow Budget")