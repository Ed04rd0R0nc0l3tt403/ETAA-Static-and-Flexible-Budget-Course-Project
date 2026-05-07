import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def compute_sales_bdg(filepath: str = "data/bic_sales_data.csv"):
    df = pd.read_csv(filepath)
    df["revenue"] = df["units_millions"] * df["unit_price_eur"]

    # --- Quarterly revenue ---
    quarterly = (
        df.groupby(["year", "quarter"])["revenue"]
        .sum()
        .reset_index()
    )
    quarterly.rename(columns={"revenue": "revenue_eur_millions"}, inplace=True)

    # --- Quarterly units sold ---
    quarterly_units = (
        df.groupby(["year", "quarter"])["units_millions"]
        .sum()
        .reset_index()
    )
    quarterly_units.rename(columns={"units_millions": "units_sold_millions"}, inplace=True)

    # --- Merge revenue and units into one quarterly DataFrame ---
    quarterly = quarterly.merge(quarterly_units, on=["year", "quarter"], how="left")

    # --- Total revenue per year ---
    yearly_revenue = (
        quarterly
        .groupby("year")["revenue_eur_millions"]
        .sum()
        .reset_index()
        .rename(columns={"revenue_eur_millions": "total_revenue_eur_millions"})
    )

    # --- Total revenue across all years ---
    total_revenue = quarterly["revenue_eur_millions"].sum()

    return quarterly, yearly_revenue, total_revenue


def export_sales_bdg_excel(
    filepath:     str = "data/bic_sales_data.csv",
    output_path:  str = "output/sales_budget.xlsx"
):
    quarterly, yearly_revenue, total_revenue = compute_sales_bdg(filepath)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Budget"

    # --- Styles ---
    header_font    = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill    = PatternFill("solid", start_color="1F4E79")
    section_fill   = PatternFill("solid", start_color="2E75B6")
    total_fill     = PatternFill("solid", start_color="D6E4F0")
    total_font     = Font(name="Arial", bold=True, color="000000")
    center         = Alignment(horizontal="center")
    thin           = Side(style="thin", color="CCCCCC")
    border         = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_cell(cell, font=None, fill=None, alignment=None, number_format=None):
        if font:           cell.font           = font
        if fill:           cell.fill           = fill
        if alignment:      cell.alignment      = alignment
        if number_format:  cell.number_format  = number_format
        cell.border = border

    # --- Title ---
    ws.merge_cells("A1:E1")
    ws["A1"] = "BIC Corporation — Sales Budget"
    ws["A1"].font      = Font(name="Arial", bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    # --- Section 1: Quarterly Revenue and Units ---
    ws.append([])
    ws.append(["QUARTERLY DETAIL", "", "", "", ""])
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    headers = ["Year", "Quarter", "Units Sold (M)", "Revenue (€M)"]
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=ws.max_row, column=col)
        style_cell(cell, font=header_font, fill=header_fill, alignment=center)

    for _, row in quarterly.iterrows():
        ws.append([
            int(row["year"]),
            row["quarter"],
            f'=C{ws.max_row + 1}',   # placeholder — written below
            None
        ])
        r = ws.max_row
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = row["quarter"]
        ws.cell(r, 3).value = round(row["units_sold_millions"], 4)
        ws.cell(r, 4).value = round(row["revenue_eur_millions"], 4)
        for col in range(1, 5):
            style_cell(ws.cell(r, col), number_format='#,##0.00;(#,##0.00);-')

    # --- Section 2: Yearly Revenue ---
    ws.append([])
    ws.append(["YEARLY REVENUE SUMMARY", "", "", ""])
    ws[f"A{ws.max_row}"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws[f"A{ws.max_row}"].fill = section_fill

    yr_headers = ["Year", "Total Revenue (€M)"]
    ws.append(yr_headers)
    for col, _ in enumerate(yr_headers, 1):
        style_cell(ws.cell(ws.max_row, col), font=header_font, fill=header_fill, alignment=center)

    for _, row in yearly_revenue.iterrows():
        r = ws.max_row + 1
        ws.cell(r, 1).value = int(row["year"])
        ws.cell(r, 2).value = round(row["total_revenue_eur_millions"], 4)
        for col in range(1, 3):
            style_cell(ws.cell(r, col), number_format='#,##0.00;(#,##0.00);-')

    # --- Total row ---
    ws.append([])
    r = ws.max_row + 1
    ws.cell(r, 1).value = "TOTAL REVENUE (ALL YEARS)"
    ws.cell(r, 2).value = round(total_revenue, 4)
    for col in range(1, 3):
        style_cell(ws.cell(r, col), font=total_font, fill=total_fill,
                   number_format='#,##0.00;(#,##0.00);-')

    # --- Column widths ---
    col_widths = [10, 12, 18, 18]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(output_path)
    print(f"Sales budget exported to {output_path}")
